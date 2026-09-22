import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def create_docx():
    doc = docx.Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base styling
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Read Markdown
    with open('Facebook_Automated_Posting_System_Report.md', 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []

    for line in lines:
        raw_line = line.rstrip('\n')

        # Code block boundary
        if raw_line.startswith('```'):
            if in_code_block:
                # End of code block
                in_code_block = False
                code_text = '\n'.join(code_lines)
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.4)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(code_text)
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(raw_line)
            continue

        # Table processing
        if '|' in raw_line and not raw_line.startswith('```'):
            parts = [p.strip() for p in raw_line.split('|')[1:-1]]
            if len(parts) > 0:
                if any(set(p) == {'-'} or set(p) == {':', '-'} for p in parts if p):
                    # Separator line, skip
                    continue
                table_rows.append(parts)
                in_table = True
                continue

        if in_table and ('|' not in raw_line or not raw_line.strip()):
            # Render accumulated table
            if table_rows:
                num_cols = max(len(r) for r in table_rows)
                table = doc.add_table(rows=len(table_rows), cols=num_cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER

                for r_idx, row_data in enumerate(table_rows):
                    for c_idx, cell_value in enumerate(row_data):
                        if c_idx < num_cols:
                            cell = table.cell(r_idx, c_idx)
                            cell.text = cell_value.replace('**', '').replace('`', '')
                            # Header styling
                            if r_idx == 0:
                                set_cell_background(cell, "003366")
                                for p in cell.paragraphs:
                                    for r in p.runs:
                                        r.font.bold = True
                                        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                            else:
                                if r_idx % 2 == 1:
                                    set_cell_background(cell, "F2F4F8")
                                else:
                                    set_cell_background(cell, "FFFFFF")
                doc.add_paragraph() # spacing
            in_table = False
            table_rows = []

        # Headings
        if raw_line.startswith('# '):
            p = doc.add_heading(level=1)
            run = p.add_run(raw_line[2:])
            run.font.name = 'Calibri'
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
        elif raw_line.startswith('## '):
            p = doc.add_heading(level=2)
            run = p.add_run(raw_line[3:])
            run.font.name = 'Calibri'
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x00, 0x55, 0x99)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
        elif raw_line.startswith('### '):
            p = doc.add_heading(level=3)
            run = p.add_run(raw_line[4:])
            run.font.name = 'Calibri'
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x00, 0x66, 0xCC)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
        elif raw_line.startswith('* ') or raw_line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            # Simple inline bold processing
            text = raw_line[2:]
            parts = text.split('**')
            for idx, part in enumerate(parts):
                run = p.add_run(part)
                if idx % 2 == 1:
                    run.bold = True
        elif raw_line.strip() == '---':
            continue
        elif raw_line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            text = raw_line
            parts = text.split('**')
            for idx, part in enumerate(parts):
                run = p.add_run(part)
                if idx % 2 == 1:
                    run.bold = True

    # Render any remaining table at end of file
    if table_rows:
        num_cols = max(len(r) for r in table_rows)
        table = doc.add_table(rows=len(table_rows), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r_idx, row_data in enumerate(table_rows):
            for c_idx, cell_value in enumerate(row_data):
                if c_idx < num_cols:
                    cell = table.cell(r_idx, c_idx)
                    cell.text = cell_value.replace('**', '').replace('`', '')
                    if r_idx == 0:
                        set_cell_background(cell, "003366")
                        for p in cell.paragraphs:
                            for r in p.runs:
                                r.font.bold = True
                                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    else:
                        if r_idx % 2 == 1:
                            set_cell_background(cell, "F2F4F8")
                        else:
                            set_cell_background(cell, "FFFFFF")

    doc.save('Facebook_Automated_Posting_System_Report.docx')
    print("[SUCCESS] Successfully generated Facebook_Automated_Posting_System_Report.docx")

create_docx()
